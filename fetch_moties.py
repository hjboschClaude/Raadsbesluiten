#!/usr/bin/env python3
"""Haal alle moties op van gemeenteraad.rotterdam.nl en sla ze op als Excel.

Voor aangenomen moties worden ook de detailpagina's opgehaald om
portefeuillehouder, stand van zaken en afdoeningsstatus te extraheren.
"""

import urllib.request
import json
import re
import time
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

BASE_URL = "https://gemeenteraad.rotterdam.nl"
LIST_ID = "a61fab39-bc62-464f-968d-db31925a66e5"
API_URL = f"{BASE_URL}/Reports/GetReportData/{LIST_ID}"
PAGE_SIZE = 100
CONCURRENT_REQUESTS = 3  # Voorzichtig met rate limiting

COLUMNS_PARAM = "&".join([
    "columns[0][data]=externalid&columns[0][name]=externalid&columns[0][searchable]=true",
    "columns[1][data]=title&columns[1][name]=title&columns[1][searchable]=true",
    "columns[2][data]=partij&columns[2][name]=partij&columns[2][searchable]=true",
    "columns[3][data]=registrationdate&columns[3][name]=registrationdate&columns[3][searchable]=true",
    "columns[4][data]=uitslag&columns[4][name]=uitslag&columns[4][searchable]=true",
    "columns[5][data]=completed&columns[5][name]=completed&columns[5][searchable]=true",
    "columns[6][data]=datecompleted&columns[6][name]=datecompleted&columns[6][searchable]=true",
    "columns[7][data]=medeondertekenaars&columns[7][name]=medeondertekenaars&columns[7][searchable]=true",
    "columns[8][data]=medeindiendepartijen&columns[8][name]=medeindiendepartijen&columns[8][searchable]=true",
    "columns[9][data]=raadslid&columns[9][name]=raadslid&columns[9][searchable]=true",
])


# --- HTTP helpers ---

def http_request(url, max_retries=4, timeout=30, data=None, headers=None):
    """HTTP request met exponential backoff retry."""
    for attempt in range(max_retries):
        try:
            req = urllib.request.Request(url, data=data, headers=headers or {})
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                return resp.read().decode("utf-8")
        except Exception as e:
            if attempt < max_retries - 1:
                wait = 2 ** (attempt + 1)
                time.sleep(wait)
            else:
                raise


# --- API data ophalen ---

def fetch_page(draw, start, length):
    params = (
        f"draw={draw}&start={start}&length={length}"
        f"&order[0][column]=3&order[0][dir]=desc"
        f"&search[value]=&search[regex]=false&{COLUMNS_PARAM}"
    )
    body = http_request(
        API_URL,
        data=params.encode("utf-8"),
        headers={
            "Content-Type": "application/x-www-form-urlencoded",
            "Referer": f"{BASE_URL}/Reports/Details/{LIST_ID}",
        },
    )
    return json.loads(body)


def fetch_all_records():
    all_records = []
    start = 0
    draw = 1
    total = None

    while True:
        result = fetch_page(draw, start, PAGE_SIZE)
        if total is None:
            total = result["recordsTotal"]
            print(f"Totaal aantal moties: {total}")
        all_records.extend(result["data"])
        print(f"  Lijst opgehaald: {len(all_records)}/{total}")
        if len(all_records) >= total or len(result["data"]) == 0:
            break
        start += PAGE_SIZE
        draw += 1

    # Normaliseer uitslag (inconsistent hoofdlettergebruik in de data)
    for r in all_records:
        u = (r.get("uitslag") or "").strip()
        if u.lower() == "aangenomen":
            r["uitslag"] = "Aangenomen"
        elif u.lower() == "verworpen":
            r["uitslag"] = "Verworpen"
        elif u.lower() == "ingetrokken":
            r["uitslag"] = "Ingetrokken"
        elif u.lower() == "aangehouden":
            r["uitslag"] = "Aangehouden"

    return all_records


# --- Detail pagina parsing ---

def get_text_field(html, label):
    """Haal tekstveld op uit dt/dd structuur."""
    pattern = r'<dt[^>]*>\s*' + re.escape(label) + r'\s*</dt>\s*<dd[^>]*>(.*?)</dd>'
    match = re.search(pattern, html, re.DOTALL)
    if not match:
        return ""
    content = match.group(1)
    # Check for pre-line spans (used for longer text)
    pre = re.search(r'<span class="pre-line">(.*?)</span>', content, re.DOTALL)
    if pre:
        text = re.sub(r'<[^>]+>', ' ', pre.group(1)).strip()
    else:
        text = re.sub(r'<[^>]+>', ' ', content).strip()
    return re.sub(r'\s+', ' ', text).strip()


def get_list_field(html, label):
    """Haal lijstveld op (ul/li) uit dt/dd structuur."""
    pattern = r'<dt[^>]*>\s*' + re.escape(label) + r'\s*</dt>\s*<dd[^>]*>(.*?)</dd>'
    match = re.search(pattern, html, re.DOTALL)
    if not match:
        return ""
    content = match.group(1)
    items = re.findall(r'<li[^>]*>(.*?)</li>', content, re.DOTALL)
    if items:
        return ", ".join(re.sub(r'<[^>]+>', '', item).strip() for item in items)
    return re.sub(r'<[^>]+>', ' ', content).strip().replace("  ", " ")


def get_checkbox_field(html, label):
    """Haal checkbox-veld op (Ja/Nee)."""
    pattern = r'<dt[^>]*>\s*' + re.escape(label) + r'\s*</dt>\s*<dd[^>]*>(.*?)</dd>'
    match = re.search(pattern, html, re.DOTALL)
    if not match:
        return ""
    content = match.group(1)
    if 'fa-check-square' in content:
        return "Ja"
    if 'fa-square' in content:
        return "Nee"
    return ""


def get_document_url(html, row_id):
    """Extraheer de PDF-document URL uit het Hoofddocument veld."""
    pattern = r'<dt[^>]*>\s*Hoofddocument\s*</dt>\s*<dd[^>]*>(.*?)</dd>'
    match = re.search(pattern, html, re.DOTALL)
    if not match:
        return ""
    doc_link = re.search(r'href="(/Reports/Document/[^"]+)"', match.group(1))
    if doc_link:
        return f"{BASE_URL}{doc_link.group(1)}"
    return ""


def fetch_item_details(record):
    """Haal detailpagina op voor een motie en extraheer statusvelden."""
    row_id = record["DT_RowId"]
    url = f"{BASE_URL}/Reports/Item/{row_id}"

    try:
        html = http_request(url)

        record["document_url"] = get_document_url(html, row_id)
        record["portefeuillehouder"] = get_list_field(html, "Portefeuillehouder")
        record["beleidsveld"] = get_list_field(html, "Beleidsveld")
        record["commissie"] = get_text_field(html, "Commissie")
        record["omschrijving"] = get_text_field(html, "Omschrijving")
        record["verwachte_datum_afdoening"] = get_text_field(html, "Verwachte datum afdoening")
        record["stand_van_zaken"] = get_text_field(html, "Stand van zaken")
        record["afgedaan"] = get_checkbox_field(html, "Afgedaan")
        record["afdoeningsvoorstel_aanwezig"] = get_checkbox_field(html, "Afdoeningsvoorstel aanwezig")
        record["toelichting"] = get_text_field(html, "Toelichting")
        record["afdoening"] = get_text_field(html, "Afdoening")
        record["detail_opgehaald"] = True

    except Exception as e:
        print(f"  FOUT bij {row_id} ({record.get('externalid','')}): {e}", file=sys.stderr)
        for key in ["document_url", "portefeuillehouder", "beleidsveld", "commissie",
                     "omschrijving", "verwachte_datum_afdoening", "stand_van_zaken",
                     "afgedaan", "afdoeningsvoorstel_aanwezig", "toelichting", "afdoening"]:
            record[key] = ""
        record["detail_opgehaald"] = False

    return record


def fetch_details_for_aangenomen(records):
    """Haal details op voor alle aangenomen moties."""
    aangenomen = [r for r in records if r.get("uitslag") == "Aangenomen"]
    total = len(aangenomen)
    print(f"  {total} aangenomen moties gevonden, details ophalen...")
    completed = 0

    with ThreadPoolExecutor(max_workers=CONCURRENT_REQUESTS) as executor:
        futures = {executor.submit(fetch_item_details, r): r for r in aangenomen}
        for future in as_completed(futures):
            completed += 1
            if completed % 100 == 0 or completed == total:
                print(f"  Details opgehaald: {completed}/{total}")

    return records


def fetch_document_url(record):
    """Haal alleen de document-URL op voor een motie."""
    row_id = record["DT_RowId"]
    url = f"{BASE_URL}/Reports/Item/{row_id}"
    try:
        html = http_request(url)
        record["document_url"] = get_document_url(html, row_id)
    except Exception as e:
        print(f"  FOUT bij document URL {row_id} ({record.get('externalid','')}): {e}", file=sys.stderr)
        record["document_url"] = ""
    return record


def fetch_document_urls_for_overige(records):
    """Haal document-URLs op voor niet-aangenomen moties."""
    overige = [r for r in records if r.get("uitslag") != "Aangenomen"]
    total = len(overige)
    print(f"  {total} overige moties, document-URLs ophalen...")
    completed = 0

    with ThreadPoolExecutor(max_workers=CONCURRENT_REQUESTS) as executor:
        futures = {executor.submit(fetch_document_url, r): r for r in overige}
        for future in as_completed(futures):
            completed += 1
            if completed % 100 == 0 or completed == total:
                print(f"  Document-URLs opgehaald: {completed}/{total}")

    return records


# --- Excel generatie ---

def make_styles():
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="00674A", end_color="00674A", fill_type="solid")
    header_fill_2 = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_alignment = Alignment(vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    return {
        "header_font": header_font,
        "header_fill": header_fill,
        "header_fill_2": header_fill_2,
        "header_alignment": header_alignment,
        "cell_alignment": cell_alignment,
        "thin_border": thin_border,
    }


def write_sheet(ws, columns, records, styles):
    """Schrijf data naar een werkblad."""
    # Headers
    for col_idx, (header, _, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.alignment = styles["header_alignment"]
        cell.border = styles["thin_border"]
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Data
    for row_idx, record in enumerate(records, 2):
        for col_idx, (_, key, _) in enumerate(columns, 1):
            value = record.get(key)
            if value is None:
                value = ""
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = styles["cell_alignment"]
            cell.border = styles["thin_border"]

    # Freeze en filter
    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(columns))
    ws.auto_filter.ref = f"A1:{last_col}{len(records) + 1}"


def create_excel(records, filename):
    wb = Workbook()
    styles = make_styles()

    # --- Blad 1: Alle moties ---
    ws1 = wb.active
    ws1.title = "Alle moties"
    alle_kolommen = [
        ("BB-nummer", "externalid", 16),
        ("Titel", "title", 60),
        ("Partij", "partij", 18),
        ("Raadslid", "raadslid", 22),
        ("Datum ingediend", "registrationdate", 16),
        ("Uitslag", "uitslag", 14),
        ("Afgedaan", "afgedaan_api", 12),
        ("Datum afgedaan", "datecompleted", 16),
        ("Medeondertekenaars", "medeondertekenaars", 30),
        ("Mede indienende partijen", "medeindiendepartijen", 25),
        ("URL", "url", 50),
        ("Document", "document_url", 50),
    ]

    # Voeg url en afgedaan_api toe
    for r in records:
        r["url"] = f"{BASE_URL}/Reports/Item/{r['DT_RowId']}"
        if r.get("uitslag") == "Aangenomen":
            r["afgedaan_api"] = "Ja" if r.get("datecompleted") else "Nee"
        else:
            r["afgedaan_api"] = ""

    write_sheet(ws1, alle_kolommen, records, styles)

    # Kleur uitslag kolom
    uitslag_col = 6  # kolom F
    groen = Font(name="Calibri", color="1B5E20", bold=True)
    rood = Font(name="Calibri", color="B71C1C", bold=True)
    for row_idx, record in enumerate(records, 2):
        cell = ws1.cell(row=row_idx, column=uitslag_col)
        if record.get("uitslag") == "Aangenomen":
            cell.font = groen
        elif record.get("uitslag") == "Verworpen":
            cell.font = rood

    # --- Blad 2: Aangenomen moties (detail) ---
    ws2 = wb.create_sheet("Aangenomen moties")
    aangenomen = [r for r in records if r.get("uitslag") == "Aangenomen"]

    detail_kolommen = [
        ("BB-nummer", "externalid", 16),
        ("Titel", "title", 55),
        ("Partij", "partij", 18),
        ("Raadslid", "raadslid", 22),
        ("Datum ingediend", "registrationdate", 16),
        ("Portefeuillehouder", "portefeuillehouder", 35),
        ("Beleidsveld", "beleidsveld", 22),
        ("Commissie", "commissie", 30),
        ("Afgedaan", "afgedaan", 10),
        ("Datum afgedaan", "datecompleted", 16),
        ("Verwachte afdoening", "verwachte_datum_afdoening", 18),
        ("Stand van zaken", "stand_van_zaken", 60),
        ("Afdoeningsvoorstel", "afdoeningsvoorstel_aanwezig", 16),
        ("Toelichting", "toelichting", 50),
        ("Afdoening", "afdoening", 50),
        ("URL", "url", 50),
        ("Document", "document_url", 50),
    ]

    # Sorteer op portefeuillehouder, dan datum
    aangenomen.sort(key=lambda r: (r.get("portefeuillehouder", "") or "", r.get("registrationdate", "")))
    write_sheet(ws2, detail_kolommen, aangenomen, styles)

    # Kleur afgedaan kolom
    afgedaan_col = 9  # kolom I
    for row_idx, record in enumerate(aangenomen, 2):
        cell = ws2.cell(row=row_idx, column=afgedaan_col)
        if record.get("afgedaan") == "Ja":
            cell.font = groen
        elif record.get("afgedaan") == "Nee":
            cell.font = rood

    # --- Blad 3: Samenvatting per portefeuillehouder ---
    ws3 = wb.create_sheet("Per portefeuillehouder")

    # Bereken statistieken per portefeuillehouder
    ph_stats = {}
    for r in aangenomen:
        ph = r.get("portefeuillehouder", "") or "(onbekend)"
        if ph not in ph_stats:
            ph_stats[ph] = {"totaal": 0, "afgedaan": 0, "open": 0, "open_moties": []}
        ph_stats[ph]["totaal"] += 1
        if r.get("afgedaan") == "Ja" or r.get("datecompleted"):
            ph_stats[ph]["afgedaan"] += 1
        else:
            ph_stats[ph]["open"] += 1
            ph_stats[ph]["open_moties"].append(r)

    # Headers
    samenvatting_headers = [
        ("Portefeuillehouder", 35),
        ("Totaal aangenomen", 18),
        ("Afgedaan", 12),
        ("Open", 12),
        ("% Afgedaan", 12),
    ]
    for col_idx, (header, width) in enumerate(samenvatting_headers, 1):
        cell = ws3.cell(row=1, column=col_idx, value=header)
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.alignment = styles["header_alignment"]
        cell.border = styles["thin_border"]
        ws3.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, (ph, stats) in enumerate(sorted(ph_stats.items()), 2):
        pct = round(stats["afgedaan"] / stats["totaal"] * 100) if stats["totaal"] else 0
        values = [ph, stats["totaal"], stats["afgedaan"], stats["open"], f"{pct}%"]
        for col_idx, val in enumerate(values, 1):
            cell = ws3.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = styles["cell_alignment"]
            cell.border = styles["thin_border"]

    ws3.freeze_panes = "A2"
    ws3.auto_filter.ref = f"A1:E{len(ph_stats) + 1}"

    wb.save(filename)
    print(f"\nExcel bestand opgeslagen: {filename}")
    print(f"  Blad 'Alle moties': {len(records)} rijen")
    print(f"  Blad 'Aangenomen moties': {len(aangenomen)} rijen (gesorteerd op portefeuillehouder)")
    print(f"  Blad 'Per portefeuillehouder': {len(ph_stats)} portefeuillehouders")


# --- Main ---

if __name__ == "__main__":
    print("Stap 1: Alle moties ophalen via API...")
    records = fetch_all_records()

    # Statistieken
    uitslagen = {}
    for r in records:
        u = r.get("uitslag") or "(geen uitslag)"
        uitslagen[u] = uitslagen.get(u, 0) + 1
    print("\n  Uitslagen:")
    for u, c in sorted(uitslagen.items(), key=lambda x: -x[1]):
        print(f"    {u}: {c}")

    print(f"\nStap 2: Details ophalen voor aangenomen moties...")
    fetch_details_for_aangenomen(records)

    # Detail statistieken
    aangenomen = [r for r in records if r.get("uitslag") == "Aangenomen"]
    opgehaald = sum(1 for r in aangenomen if r.get("detail_opgehaald"))
    afgedaan = sum(1 for r in aangenomen if r.get("afgedaan") == "Ja" or r.get("datecompleted"))
    print(f"\n  Details succesvol opgehaald: {opgehaald}/{len(aangenomen)}")
    print(f"  Afgedaan: {afgedaan}")
    print(f"  Nog open: {len(aangenomen) - afgedaan}")

    print(f"\nStap 3: Document-URLs ophalen voor overige moties...")
    fetch_document_urls_for_overige(records)

    print("\nStap 4: Excel genereren...")
    create_excel(records, "moties.xlsx")
