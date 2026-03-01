#!/usr/bin/env python3
"""Haal alle schriftelijke vragen op van gemeenteraad.rotterdam.nl en sla ze op als Excel.

Geeft aan hoeveel schriftelijke vragen zijn gedownload (opgehaald) van het portaal.
"""

import urllib.request
import json
import re
import time
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

BASE_URL = "https://gemeenteraad.rotterdam.nl"
LIST_ID = "da9b533f-5f24-4f51-8567-19fe410f15d4"
API_URL = f"{BASE_URL}/Reports/GetReportData/{LIST_ID}"
PAGE_SIZE = 100
CONCURRENT_REQUESTS = 3  # Voorzichtig met rate limiting

COLUMNS_PARAM = "&".join([
    "columns[0][data]=externalid&columns[0][name]=externalid&columns[0][searchable]=true",
    "columns[1][data]=title&columns[1][name]=title&columns[1][searchable]=true",
    "columns[2][data]=partij&columns[2][name]=partij&columns[2][searchable]=true",
    "columns[3][data]=registrationdate&columns[3][name]=registrationdate&columns[3][searchable]=true",
    "columns[4][data]=datecompleted&columns[4][name]=datecompleted&columns[4][searchable]=true",
    "columns[5][data]=beleidsveld&columns[5][name]=beleidsveld&columns[5][searchable]=true",
    "columns[6][data]=raadslid&columns[6][name]=raadslid&columns[6][searchable]=true",
    "columns[7][data]=medeondertekenaars&columns[7][name]=medeondertekenaars&columns[7][searchable]=true",
    "columns[8][data]=medeindiendepartijen&columns[8][name]=medeindiendepartijen&columns[8][searchable]=true",
])

EXCEL_COLUMNS = [
    ("BB-nummer", "externalid", 16),
    ("Titel", "title", 70),
    ("Partij", "partij", 18),
    ("Raadslid", "raadslid", 22),
    ("Datum ingediend", "registrationdate", 16),
    ("Datum beantwoord", "datecompleted", 16),
    ("Beleidsveld", "beleidsveld", 25),
    ("Portefeuillehouder", "portefeuillehouder", 35),
    ("Commissie", "commissie", 35),
    ("Medeondertekenaars", "medeondertekenaars", 30),
    ("Mede indienende partijen", "medeindiendepartijen", 25),
    ("Verwachte datum afdoening", "verwachte_datum_afdoening", 20),
    ("Stand van zaken", "stand_van_zaken", 50),
    ("Hoofddocument", "hoofddocument_naam", 60),
    ("Hoofddocument URL", "hoofddocument_url", 50),
    ("URL", "url", 50),
]


def http_request(url, max_retries=4, timeout=30, data=None, headers=None):
    """HTTP request met exponential backoff retry."""
    for attempt in range(max_retries):
        try:
            req = urllib.request.Request(url, data=data, headers=headers or {})
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                return resp.read().decode("utf-8")
        except Exception as e:
            if attempt < max_retries - 1:
                wait = 2 ** (attempt + 1)
                time.sleep(wait)
            else:
                raise


def fetch_page(draw, start, length):
    params = (
        f"draw={draw}&start={start}&length={length}"
        f"&order[0][column]=3&order[0][dir]=desc"
        f"&search[value]=&search[regex]=false&{COLUMNS_PARAM}"
    )
    body = http_request(
        API_URL,
        data=params.encode("utf-8"),
        headers={
            "Content-Type": "application/x-www-form-urlencoded",
            "Referer": f"{BASE_URL}/Reports/Details/{LIST_ID}",
        },
    )
    return json.loads(body)


def fetch_all_records():
    all_records = []
    start = 0
    draw = 1
    total = None

    while True:
        result = fetch_page(draw, start, PAGE_SIZE)
        if total is None:
            total = result["recordsTotal"]
            print(f"Totaal aantal schriftelijke vragen in het portaal: {total}")

        all_records.extend(result["data"])
        print(f"  Lijst opgehaald: {len(all_records)}/{total}")

        if len(all_records) >= total or len(result["data"]) == 0:
            break

        start += PAGE_SIZE
        draw += 1

    return all_records


def get_text_field(html, label):
    """Haal tekstveld op uit dt/dd structuur."""
    pattern = r'<dt[^>]*>\s*' + re.escape(label) + r'\s*</dt>\s*<dd[^>]*>(.*?)</dd>'
    match = re.search(pattern, html, re.DOTALL)
    if not match:
        return ""
    content = match.group(1)
    pre = re.search(r'<span class="pre-line">(.*?)</span>', content, re.DOTALL)
    if pre:
        text = re.sub(r'<[^>]+>', ' ', pre.group(1)).strip()
    else:
        text = re.sub(r'<[^>]+>', ' ', content).strip()
    return re.sub(r'\s+', ' ', text).strip()


def get_list_field(html, label):
    """Haal lijstveld op (ul/li) uit dt/dd structuur."""
    pattern = r'<dt[^>]*>\s*' + re.escape(label) + r'\s*</dt>\s*<dd[^>]*>(.*?)</dd>'
    match = re.search(pattern, html, re.DOTALL)
    if not match:
        return ""
    content = match.group(1)
    items = re.findall(r'<li[^>]*>(.*?)</li>', content, re.DOTALL)
    if items:
        return ", ".join(re.sub(r'<[^>]+>', '', item).strip() for item in items)
    return re.sub(r'<[^>]+>', ' ', content).strip().replace("  ", " ")


def get_document_url(html):
    """Extraheer de hoofddocument URL en naam."""
    pattern = r'<dt[^>]*>\s*Hoofddocument\s*</dt>\s*<dd[^>]*>(.*?)</dd>'
    match = re.search(pattern, html, re.DOTALL)
    if not match:
        return "", ""
    content = match.group(1)
    doc_link = re.search(
        r'href="(/Reports/Document/[^"]+)"[^>]*data-document-id="[^"]*"[^>]*>.*?<span[^>]*>[^<]*</span>\s*(.*?)\s*<span class="badge',
        content,
        re.DOTALL,
    )
    if doc_link:
        url = f"{BASE_URL}{doc_link.group(1)}"
        name = re.sub(r'<[^>]+>', '', doc_link.group(2)).strip()
        return url, name
    # Fallback: only URL
    link = re.search(r'href="(/Reports/Document/[^"]+)"', content)
    if link:
        return f"{BASE_URL}{link.group(1)}", ""
    return "", ""


def fetch_item_details(record):
    """Haal detailpagina op en extraheer extra velden."""
    row_id = record["DT_RowId"]
    url = f"{BASE_URL}/Reports/Item/{row_id}"

    try:
        html = http_request(url)

        doc_url, doc_naam = get_document_url(html)
        record["hoofddocument_url"] = doc_url
        record["hoofddocument_naam"] = doc_naam
        record["portefeuillehouder"] = get_list_field(html, "Portefeuillehouder")
        record["commissie"] = get_text_field(html, "Commissie")
        record["verwachte_datum_afdoening"] = get_text_field(html, "Verwachte datum afdoening")
        record["stand_van_zaken"] = get_text_field(html, "Stand van zaken")
        if not record.get("medeondertekenaars"):
            record["medeondertekenaars"] = get_text_field(html, "Medeondertekenaars")
        if not record.get("medeindiendepartijen"):
            record["medeindiendepartijen"] = get_text_field(html, "Mede indienende partijen")
        record["detail_opgehaald"] = True

    except Exception as e:
        print(f"  FOUT bij {row_id} ({record.get('externalid', '')}): {e}", file=sys.stderr)
        for key in ["hoofddocument_url", "hoofddocument_naam", "portefeuillehouder",
                    "commissie", "verwachte_datum_afdoening", "stand_van_zaken"]:
            record.setdefault(key, "")
        record["detail_opgehaald"] = False

    return record


def fetch_all_details(records):
    """Haal detailpagina's op voor alle records."""
    total = len(records)
    completed = 0
    download_count = 0

    with ThreadPoolExecutor(max_workers=CONCURRENT_REQUESTS) as executor:
        futures = {executor.submit(fetch_item_details, r): r for r in records}
        for future in as_completed(futures):
            result = future.result()
            completed += 1
            if result.get("detail_opgehaald"):
                download_count += 1
            if completed % 200 == 0 or completed == total:
                print(f"  Details opgehaald: {completed}/{total} "
                      f"(succesvol: {download_count})")

    return records, download_count


def make_styles():
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1A4B7A", end_color="1A4B7A", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_alignment = Alignment(vertical="top", wrap_text=True)
    url_font = Font(name="Calibri", color="0563C1", underline="single", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    return {
        "header_font": header_font,
        "header_fill": header_fill,
        "header_alignment": header_alignment,
        "cell_alignment": cell_alignment,
        "url_font": url_font,
        "thin_border": thin_border,
    }


def write_sheet(ws, columns, records, styles):
    """Schrijf data naar een werkblad."""
    for col_idx, (header, _, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.alignment = styles["header_alignment"]
        cell.border = styles["thin_border"]
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, record in enumerate(records, 2):
        for col_idx, (_, key, _) in enumerate(columns, 1):
            value = record.get(key) or ""
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = styles["cell_alignment"]
            cell.border = styles["thin_border"]
            if key == "hoofddocument_url" and value:
                cell.hyperlink = value
                cell.font = styles["url_font"]
            elif key == "url" and value:
                cell.hyperlink = value
                cell.font = styles["url_font"]

    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(columns))
    ws.auto_filter.ref = f"A1:{last_col}{len(records) + 1}"


def create_excel(records, filename, download_count):
    wb = Workbook()
    styles = make_styles()

    # Blad 1: Alle schriftelijke vragen
    ws1 = wb.active
    ws1.title = "Schriftelijke vragen"

    for r in records:
        r["url"] = f"{BASE_URL}/Reports/Item/{r['DT_RowId']}"

    write_sheet(ws1, EXCEL_COLUMNS, records, styles)

    # Blad 2: Samenvatting per partij
    ws2 = wb.create_sheet("Per partij")
    partij_stats = {}
    for r in records:
        partij = (r.get("partij") or "").strip() or "(onbekend)"
        if partij not in partij_stats:
            partij_stats[partij] = {"totaal": 0, "beantwoord": 0}
        partij_stats[partij]["totaal"] += 1
        if r.get("datecompleted"):
            partij_stats[partij]["beantwoord"] += 1

    samenvatting_headers = [
        ("Partij", 30),
        ("Aantal vragen", 15),
        ("Beantwoord", 15),
        ("Nog open", 12),
        ("% Beantwoord", 14),
    ]
    for col_idx, (header, width) in enumerate(samenvatting_headers, 1):
        cell = ws2.cell(row=1, column=col_idx, value=header)
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.alignment = styles["header_alignment"]
        cell.border = styles["thin_border"]
        ws2.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, (partij, stats) in enumerate(
        sorted(partij_stats.items(), key=lambda x: -x[1]["totaal"]), 2
    ):
        pct = round(stats["beantwoord"] / stats["totaal"] * 100) if stats["totaal"] else 0
        open_count = stats["totaal"] - stats["beantwoord"]
        values = [partij, stats["totaal"], stats["beantwoord"], open_count, f"{pct}%"]
        for col_idx, val in enumerate(values, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = styles["cell_alignment"]
            cell.border = styles["thin_border"]

    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:E{len(partij_stats) + 1}"

    wb.save(filename)

    beantwoord = sum(1 for r in records if r.get("datecompleted"))
    open_count = len(records) - beantwoord
    print(f"\nExcel bestand opgeslagen: {filename}")
    print(f"  Blad 'Schriftelijke vragen': {len(records)} rijen")
    print(f"  Blad 'Per partij': {len(partij_stats)} partijen")
    print(f"\n=== Samenvatting download ===")
    print(f"  Totaal opgehaald (gedownload): {download_count} schriftelijke vragen")
    print(f"  Beantwoord:                    {beantwoord}")
    print(f"  Nog open (onbeantwoord):       {open_count}")


if __name__ == "__main__":
    print("Stap 1: Schriftelijke vragen ophalen via API...")
    records = fetch_all_records()

    print(f"\nStap 2: Detailpagina's ophalen voor {len(records)} schriftelijke vragen...")
    records, download_count = fetch_all_details(records)

    fouten = sum(1 for r in records if not r.get("detail_opgehaald"))
    if fouten:
        print(f"  {fouten} vragen konden niet opgehaald worden")

    print("\nStap 3: Excel genereren...")
    create_excel(records, "data/schriftelijke_vragen.xlsx", download_count)
