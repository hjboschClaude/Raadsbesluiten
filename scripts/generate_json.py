#!/usr/bin/env python3
"""Converteert data/schriftelijke_vragen.xlsx naar data/vragen.json.

Bereken ook de doorlooptijd (in kalenderdagen) per vraag.
"""

import json
import os
from datetime import datetime
from openpyxl import load_workbook

EXCEL_FILE = "data/schriftelijke_vragen.xlsx"
JSON_FILE = "data/vragen.json"
VANDAAG = datetime.today().date()


def parse_date(value):
    """Parseer datum van string of datetime object naar ISO-formaat (YYYY-MM-DD)."""
    if not value:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if hasattr(value, "strftime"):          # date object
        return value.strftime("%Y-%m-%d")
    if isinstance(value, str):
        s = value.strip()
        for fmt in ["%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y"]:
            try:
                return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
            except ValueError:
                continue
    return None


def calc_doorlooptijd(datum_ingediend, datum_beantwoord):
    """Bereken doorlooptijd in kalenderdagen.

    - Beantwoord: datum_beantwoord − datum_ingediend
    - Nog open:   vandaag − datum_ingediend  (lopende teller)
    Retourneert None als datum_ingediend ontbreekt.
    """
    if not datum_ingediend:
        return None
    try:
        d_in = datetime.strptime(datum_ingediend, "%Y-%m-%d").date()
        d_uit = (
            datetime.strptime(datum_beantwoord, "%Y-%m-%d").date()
            if datum_beantwoord
            else VANDAAG
        )
        return (d_uit - d_in).days
    except Exception:
        return None


def kwartaal(datum_str):
    """Geef kwartaalnotatie, bv. '2024-Q2'."""
    if not datum_str:
        return None
    try:
        d = datetime.strptime(datum_str, "%Y-%m-%d")
        q = (d.month - 1) // 3 + 1
        return f"{d.year}-Q{q}"
    except Exception:
        return None


def main():
    os.makedirs(os.path.dirname(JSON_FILE), exist_ok=True)

    wb = load_workbook(EXCEL_FILE, read_only=True, data_only=True)
    ws = wb.active

    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {h: i for i, h in enumerate(headers) if h}
    print(f"Kolommen gevonden ({len(headers)}): {headers}")

    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue

        def g(name):
            idx = col.get(name)
            if idx is None or idx >= len(row):
                return ""
            v = row[idx]
            if v is None:
                return ""
            if isinstance(v, datetime) or hasattr(v, "strftime"):
                return v
            return str(v).strip()

        datum_ingediend = parse_date(g("Datum ingediend"))
        datum_beantwoord = parse_date(g("Datum beantwoord"))
        beantwoord = bool(datum_beantwoord)
        doorlooptijd = calc_doorlooptijd(datum_ingediend, datum_beantwoord)

        jaar = None
        if datum_ingediend:
            try:
                jaar = int(datum_ingediend[:4])
            except Exception:
                pass

        records.append({
            "id": g("BB-nummer"),
            "titel": g("Titel"),
            "partij": g("Partij"),
            "raadslid": g("Raadslid"),
            "datum_ingediend": datum_ingediend,
            "datum_beantwoord": datum_beantwoord,
            "doorlooptijd_dagen": doorlooptijd,
            "beantwoord": beantwoord,
            "jaar": jaar,
            "kwartaal": kwartaal(datum_ingediend),
            "beleidsveld": g("Beleidsveld").strip(),
            "portefeuillehouder": g("Portefeuillehouder"),
            "commissie": g("Commissie"),
            "medeondertekenaars": g("Medeondertekenaars"),
            "medeindiendepartijen": g("Mede indienende partijen"),
            "verwachte_datum_afdoening": parse_date(g("Verwachte datum afdoening")),
            "stand_van_zaken": g("Stand van zaken"),
            "doc_naam": g("Hoofddocument"),
            "doc_url": g("Hoofddocument URL"),
            "portaal_url": g("URL"),
        })

    # Sorteer op datum ingediend aflopend (nieuwste eerst)
    records.sort(key=lambda r: r["datum_ingediend"] or "", reverse=True)

    with open(JSON_FILE, "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, separators=(",", ":"))

    beantwoord_n = sum(1 for r in records if r["beantwoord"])
    open_n = len(records) - beantwoord_n
    dl_waarden = [r["doorlooptijd_dagen"] for r in records
                  if r["beantwoord"] and r["doorlooptijd_dagen"] is not None]
    gem_dl = round(sum(dl_waarden) / len(dl_waarden)) if dl_waarden else None

    print(f"\nJSON opgeslagen: {JSON_FILE}")
    print(f"  Totaal vragen:        {len(records)}")
    print(f"  Beantwoord:           {beantwoord_n}")
    print(f"  Nog open:             {open_n}")
    if gem_dl is not None:
        print(f"  Gem. doorlooptijd:    {gem_dl} dagen (alleen beantwoorde vragen)")
    wb.close()


if __name__ == "__main__":
    main()
