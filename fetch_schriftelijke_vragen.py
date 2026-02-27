#!/usr/bin/env python3
"""Haal schriftelijke vragen op van gemeenteraad.rotterdam.nl en sla ze op als Excel.

Per vraag worden opgehaald:
  - Hoofddocument (de vraag zelf)
  - Bijlagen bij de vraag
  - Tussenbericht(en) via 'Relatie met' → Brieven B&W
  - Definitieve beantwoording via 'Relatie met' → Brieven B&W
"""

import urllib.request
import json
import re
import time
import sys
import html as html_module
from concurrent.futures import ThreadPoolExecutor, as_completed
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import cache_utils

BASE_URL = "https://gemeenteraad.rotterdam.nl"
LIST_ID = "da9b533f-5f24-4f51-8567-19fe410f15d4"
API_URL = f"{BASE_URL}/Reports/GetReportData/{LIST_ID}"
PAGE_SIZE = 100
MAX_RECORDS = 0     # 0 = geen limiet (volledig)
CONCURRENT_REQUESTS = 5

CACHE_NAME = "schriftelijke_vragen"
HASH_FIELDS = [
    "externalid", "title", "partij", "registrationdate",
    "datecompleted", "beleidsveld", "raadslid",
]

COLUMNS_PARAM = (
    "columns[0][data]=externalid&columns[0][name]=externalid&columns[0][searchable]=true&"
    "columns[1][data]=title&columns[1][name]=title&columns[1][searchable]=false&"
    "columns[2][data]=partij&columns[2][name]=partij&columns[2][searchable]=true&"
    "columns[3][data]=registrationdate&columns[3][name]=registrationdate&columns[3][searchable]=true&"
    "columns[4][data]=datecompleted&columns[4][name]=datecompleted&columns[4][searchable]=true&"
    "columns[5][data]=beleidsveld&columns[5][name]=beleidsveld&columns[5][searchable]=true&"
    "columns[6][data]=raadslid&columns[6][name]=raadslid&columns[6][searchable]=true"
)

EXCEL_COLUMNS = [
    ("BB-nummer",            "externalid",            18),
    ("Titel",                "title",                 80),
    ("Raadslid",             "raadslid",              30),
    ("Partij",               "partij",                15),
    ("Beleidsveld",          "beleidsveld",           25),
    ("Datum ingediend",      "registrationdate",      18),
    ("Datum afgedaan",       "datecompleted",         18),
    ("Hoofddocument",        "hoofddocument_naam",    80),
    ("Hoofddocument URL",    "hoofddocument_url",     50),
    ("Tussenbericht(en)",    "tussenberichten_tekst", 80),
    ("Tussenbericht URL(s)", "tussenberichten_urls",  50),
    ("Beantwoording",        "beantwoording_naam",    80),
    ("Beantwoording URL",    "beantwoording_url",     50),
    ("Bijlagen",             "bijlagen_tekst",        100),
    ("Bijlage URLs",         "bijlagen_urls",         100),
]


def http_get_with_retry(url, max_retries=4, timeout=30, data=None, headers=None):
    """HTTP request met exponential backoff retry."""
    for attempt in range(max_retries):
        try:
            req = urllib.request.Request(url, data=data, headers=headers or {})
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                return resp.read().decode("utf-8")
        except Exception as e:
            if attempt < max_retries - 1:
                time.sleep(2 ** (attempt + 1))
            else:
                raise


def fetch_page(draw, start, length):
    params = (
        f"draw={draw}&start={start}&length={length}"
        f"&order[0][column]=3&order[0][dir]=desc"
        f"&search[value]=&search[regex]=false"
        f"&{COLUMNS_PARAM}"
    )
    body = http_get_with_retry(
        API_URL,
        data=params.encode("utf-8"),
        headers={
            "Content-Type": "application/x-www-form-urlencoded",
            "Referer": f"{BASE_URL}/Reports/Details/{LIST_ID}",
        },
    )
    return json.loads(body)


def fetch_all_records():
    all_records = []
    start = 0
    draw = 1
    total = None

    while True:
        result = fetch_page(draw, start, PAGE_SIZE)
        if total is None:
            total = result["recordsTotal"]
            print(f"Totaal aantal schriftelijke vragen: {total}")
            if MAX_RECORDS:
                print(f"  (testmodus: maximaal {MAX_RECORDS} records)")

        records = result["data"]
        all_records.extend(records)

        if MAX_RECORDS and len(all_records) >= MAX_RECORDS:
            all_records = all_records[:MAX_RECORDS]
            print(f"  Lijst opgehaald: {len(all_records)}/{total} (testlimiet bereikt)")
            break

        print(f"  Lijst opgehaald: {len(all_records)}/{total}")

        if len(all_records) >= total or len(records) == 0:
            break

        start += PAGE_SIZE
        draw += 1

    return all_records


def parse_documents(html, section_label):
    """Parse documenten uit een dt/dd sectie van de HTML."""
    documents = []
    pattern = (
        r'<dt[^>]*>\s*' + re.escape(section_label) + r'\s*</dt>\s*<dd[^>]*>(.*?)</dd>'
    )
    match = re.search(pattern, html, re.DOTALL)
    if not match:
        return documents

    content = match.group(1)
    for doc_match in re.finditer(
        r'href="([^"]+)"[^>]*data-document-id="([^"]+)"[^>]*>\s*'
        r'<span class="icon\s+(\w+)"[^>]*>[^<]*</span>\s*'
        r'(.*?)\s*<span class="badge[^"]*">(.*?)</span>',
        content,
        re.DOTALL,
    ):
        url, doc_id, file_type, name, size = doc_match.groups()
        name = re.sub(r"<[^>]+>", "", name).strip()
        name = re.sub(r"\s+", " ", name)
        documents.append({
            "naam": name,
            "url": BASE_URL + url,
            "document_id": doc_id,
            "type": file_type,
            "grootte": size.strip(),
        })

    return documents


def parse_relaties(html):
    """Parse 'Relatie met' sectie: geeft lijst van Brieven B&W links terug.

    Elk item heeft:
      - item_url: URL naar het iBabs-item
      - titel: tekst na de type-label
      - is_tussenbericht: True als titel begint met 'Tussenbericht'
    """
    relaties = []
    rel_m = re.search(
        r'<dt[^>]*>\s*Relatie met\s*</dt>\s*<dd[^>]*>(.*?)</dd>', html, re.DOTALL
    )
    if not rel_m:
        return relaties

    for link_m in re.finditer(
        r'href="(/Reports/Item/[^"]+)"[^>]*>\s*<span[^>]*>(.*?)</span>\s*(.*?)</a>',
        rel_m.group(1),
        re.DOTALL,
    ):
        href, type_span, title_raw = link_m.groups()
        type_label = html_module.unescape(re.sub(r"<[^>]+>", "", type_span).strip().rstrip(":"))
        title = html_module.unescape(re.sub(r"<[^>]+>", "", title_raw).strip())
        title = re.sub(r"\s+", " ", title)

        if type_label == "Brieven B&W":
            relaties.append({
                "item_url": BASE_URL + href,
                "titel": title,
                "is_tussenbericht": title.lower().startswith("tussenbericht"),
            })

    return relaties


def fetch_relatie_document(relatie):
    """Haal het hoofddocument op van een gerelateerd Brieven B&W item."""
    try:
        html = http_get_with_retry(relatie["item_url"])
        docs = parse_documents(html, "Hoofddocument")
        if docs:
            relatie["document_naam"] = docs[0]["naam"]
            relatie["document_url"] = docs[0]["url"]
        else:
            relatie["document_naam"] = relatie["titel"]
            relatie["document_url"] = ""
    except Exception as e:
        relatie["document_naam"] = f"FOUT: {e}"
        relatie["document_url"] = ""
    return relatie


def fetch_item_documents(record):
    """Haal de detailpagina op en parse alle documenten (vraag + bijlagen + relaties)."""
    row_id = record["DT_RowId"]
    url = f"{BASE_URL}/Reports/Item/{row_id}"

    for attempt in range(4):
        try:
            html = http_get_with_retry(url)

            # Hoofddocument van de vraag zelf
            hoofddocumenten = parse_documents(html, "Hoofddocument")
            if hoofddocumenten:
                record["hoofddocument_naam"] = hoofddocumenten[0]["naam"]
                record["hoofddocument_url"] = hoofddocumenten[0]["url"]
            else:
                record["hoofddocument_naam"] = ""
                record["hoofddocument_url"] = ""

            # Bijlagen bij de vraag zelf
            bijlagen = parse_documents(html, "Bijlagen")
            record["aantal_bijlagen"] = len(bijlagen)
            record["bijlagen_tekst"] = "\n".join(
                f"{b['naam']} ({b['grootte']})" for b in bijlagen
            )
            record["bijlagen_urls"] = "\n".join(b["url"] for b in bijlagen)

            # Gerelateerde documenten (tussenberichten + beantwoording)
            relaties = parse_relaties(html)
            for rel in relaties:
                fetch_relatie_document(rel)

            tussenberichten = [r for r in relaties if r["is_tussenbericht"]]
            beantwoordingen = [r for r in relaties if not r["is_tussenbericht"]]

            record["tussenberichten_tekst"] = "\n".join(
                r.get("document_naam", r["titel"]) for r in tussenberichten
            )
            record["tussenberichten_urls"] = "\n".join(
                r.get("document_url", "") for r in tussenberichten
            )

            if beantwoordingen:
                record["beantwoording_naam"] = beantwoordingen[0].get(
                    "document_naam", beantwoordingen[0]["titel"]
                )
                record["beantwoording_url"] = beantwoordingen[0].get("document_url", "")
            else:
                record["beantwoording_naam"] = ""
                record["beantwoording_url"] = ""

            return record

        except Exception as e:
            if attempt < 3:
                time.sleep(2 ** (attempt + 1))
            else:
                print(f"  FOUT bij {row_id}: {e}", file=sys.stderr)
                for key in [
                    "hoofddocument_naam", "hoofddocument_url",
                    "tussenberichten_tekst", "tussenberichten_urls",
                    "beantwoording_naam", "beantwoording_url",
                    "bijlagen_tekst", "bijlagen_urls",
                ]:
                    record.setdefault(key, "")
                record["aantal_bijlagen"] = ""
                return record


def fetch_all_documents(records):
    """Haal documenten op voor alle records met parallelle requests."""
    total = len(records)
    completed = 0

    with ThreadPoolExecutor(max_workers=CONCURRENT_REQUESTS) as executor:
        futures = {
            executor.submit(fetch_item_documents, record): record
            for record in records
        }
        for future in as_completed(futures):
            completed += 1
            if completed % 25 == 0 or completed == total:
                print(f"  Documenten opgehaald: {completed}/{total}")

    return records


def create_excel(records, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = "Schriftelijke vragen"

    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="00674A", end_color="00674A", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_alignment = Alignment(vertical="top", wrap_text=True)
    url_font = Font(name="Calibri", color="0563C1", underline="single", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Headers
    for col_idx, (header, _, width) in enumerate(EXCEL_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        ws.column_dimensions[cell.column_letter].width = width

    # Data
    url_keys = {"hoofddocument_url", "beantwoording_url", "tussenberichten_urls", "bijlagen_urls"}
    for row_idx, record in enumerate(records, 2):
        for col_idx, (_, key, _) in enumerate(EXCEL_COLUMNS, 1):
            value = record.get(key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = cell_alignment
            cell.border = thin_border

            if key in url_keys and value:
                # Enkelvoudige URL: maak klikbare hyperlink
                if key in ("hoofddocument_url", "beantwoording_url") and "\n" not in str(value):
                    cell.hyperlink = value
                cell.font = url_font

    ws.freeze_panes = "A2"

    # Auto-filter op alle kolommen
    from openpyxl.utils import get_column_letter
    last_col = get_column_letter(len(EXCEL_COLUMNS))
    ws.auto_filter.ref = f"A1:{last_col}{len(records) + 1}"

    wb.save(filename)
    print(f"\nExcel bestand opgeslagen: {filename} ({len(records)} rijen)")


if __name__ == "__main__":
    print("Stap 1: Schriftelijke vragen ophalen via API...")
    api_records = fetch_all_records()

    print("\nStap 2: Vergelijken met cache...")
    cache = cache_utils.load_cache(CACHE_NAME)
    to_fetch, unchanged, stats = cache_utils.find_changes(api_records, cache, HASH_FIELDS)
    print(f"  Nieuw: {stats['nieuw']}, Gewijzigd: {stats['gewijzigd']}, Ongewijzigd: {stats['ongewijzigd']}")

    if to_fetch:
        print(f"\nStap 3: Documenten ophalen van {len(to_fetch)} detailpagina's...")
        fetch_all_documents(to_fetch)
    else:
        print("\nStap 3: Geen nieuwe of gewijzigde records, detailpagina's overgeslagen.")

    records = cache_utils.restore_order(api_records, to_fetch, unchanged)

    print("\nStap 4: Cache bijwerken...")
    cache_utils.save_cache(CACHE_NAME, records, HASH_FIELDS)

    # Statistieken
    fouten = sum(1 for r in records if not isinstance(r.get("aantal_bijlagen"), int))
    met_bijlagen = sum(
        1 for r in records
        if isinstance(r.get("aantal_bijlagen"), int) and r["aantal_bijlagen"] > 0
    )
    met_beantwoording = sum(1 for r in records if r.get("beantwoording_url"))
    met_tussenbericht = sum(1 for r in records if r.get("tussenberichten_urls"))

    if fouten:
        print(f"  {fouten} vragen konden niet opgehaald worden (fout)")
    print(f"  {met_bijlagen} vragen hebben bijlagen bij de vraag")
    print(f"  {met_tussenbericht} vragen hebben tussenbericht(en)")
    print(f"  {met_beantwoording} vragen hebben een definitieve beantwoording")

    print("\nStap 5: Excel genereren...")
    create_excel(records, "schriftelijke_vragen.xlsx")
